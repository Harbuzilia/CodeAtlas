from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

try:
    from docx.oxml.ns import qn
except Exception:  # pragma: no cover - optional for some docx versions
    qn = None

TEMPLATE_FILENAME = "RECH.docx"
DEFAULT_INPUT_FILENAME = "RECH_DIPLOM_v2.docx"
DEFAULT_OUTPUT_FILENAME = "RECH_DIPLOM_v2_output.docx"
FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(14)
STANDARD_SPACE = Pt(5)
LINE_SPACING = 1.0

NUMBER_PREFIX_RE = re.compile(r"^\s*(\d+)\.\s+")
QUESTION_RE = re.compile(r"^(Вопрос|Ответ):\s*(.*)$")


def set_font_style(doc: Document) -> None:
    """Устанавливает шрифт для основных стилей документа."""
    for style_name in ("Normal", "Normal (Web)", "List Paragraph", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        font = style.font
        font.name = FONT_NAME
        font.size = FONT_SIZE
        if qn is not None:
            try:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
            except Exception:
                pass


def clear_document(doc: Document) -> None:
    """Очищает содержимое документа, сохраняя параметры раздела."""
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def extract_numbering_templates(doc: Document) -> list:
    """Сохраняет шаблоны нумерации из документа-образца."""
    num_pr_by_id: dict[int | None, object] = {}
    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is None or ppr.numPr is None:
            continue
        num_pr = ppr.numPr
        num_id = num_pr.numId.val if num_pr.numId is not None else None
        if num_id not in num_pr_by_id:
            num_pr_by_id[num_id] = deepcopy(num_pr)
    return list(num_pr_by_id.values())


def apply_numbering(paragraph, num_pr) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if ppr.numPr is not None:
        ppr.remove(ppr.numPr)
    ppr.append(deepcopy(num_pr))


def apply_paragraph_format(paragraph, alignment, spacing) -> None:
    paragraph.alignment = alignment
    pf = paragraph.paragraph_format
    pf.line_spacing = LINE_SPACING
    if spacing is None:
        pf.space_before = None
        pf.space_after = None
    else:
        pf.space_before = spacing
        pf.space_after = spacing


def add_run(paragraph, text: str, bold: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    return run


def add_paragraph_with_runs(doc, runs, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=STANDARD_SPACE):
    paragraph = doc.add_paragraph()
    apply_paragraph_format(paragraph, alignment, spacing)
    for text, bold in runs:
        if not text:
            continue
        add_run(paragraph, text, bold=bold)
    return paragraph


def add_runs_from_source(doc, source_paragraph, alignment, spacing):
    paragraph = doc.add_paragraph()
    apply_paragraph_format(paragraph, alignment, spacing)
    for run in source_paragraph.runs:
        if not run.text:
            continue
        add_run(paragraph, run.text, bold=bool(run.bold))
    return paragraph


def normalize_slide_label(text: str) -> str:
    label = text.strip()
    if label.startswith("(") and label.endswith(")"):
        label = label[1:-1].strip()
    if label.startswith("[") and label.endswith("]"):
        label = label[1:-1].strip()
    return f"[{label}]"


def is_slide_heading(text: str) -> bool:
    stripped = text.strip()
    return "Слайд" in stripped and stripped.startswith(("(", "["))


def is_list_item(paragraph) -> bool:
    text = paragraph.text.strip()
    if not text:
        return False
    if paragraph.style is not None and paragraph.style.name.startswith("List"):
        return True
    return bool(NUMBER_PREFIX_RE.match(text))


def strip_number_prefix(text: str) -> str:
    return NUMBER_PREFIX_RE.sub("", text.strip())


def is_all_bold(paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False
    return all(run.bold for run in runs)


def split_schema_or_poster(text: str) -> tuple[str, str] | None:
    for prefix in ("Схема ", "Плакат "):
        if text.startswith(prefix):
            if ". " in text:
                first, rest = text.split(". ", 1)
                return f"{first}.", rest
            return text, ""
    return None


def next_non_empty(paragraphs, start_index: int):
    for idx in range(start_index + 1, len(paragraphs)):
        text = paragraphs[idx].text.strip()
        if text:
            return paragraphs[idx]
    return None


def format_document(source_path: Path, output_path: Path, template_path: Path) -> Path:
    if not template_path.exists():
        raise FileNotFoundError(f"Не найден файл шаблона: {template_path}")
    if not source_path.exists():
        raise FileNotFoundError(f"Не найден исходный документ: {source_path}")

    template = Document(str(template_path))
    numbering_templates = extract_numbering_templates(template)
    if not numbering_templates:
        raise ValueError("В шаблоне не найдены стили нумерации для списков.")

    clear_document(template)
    set_font_style(template)

    source_doc = Document(str(source_path))

    list_group_index = -1
    current_num_pr = None
    in_list = False
    first_paragraph = True

    for idx, paragraph in enumerate(source_doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            in_list = False
            continue

        if is_list_item(paragraph):
            if not in_list:
                list_group_index += 1
                if list_group_index < len(numbering_templates):
                    current_num_pr = numbering_templates[list_group_index]
                else:
                    current_num_pr = numbering_templates[-1]
                in_list = True
            spacing = None if list_group_index == 0 else STANDARD_SPACE
            item_text = strip_number_prefix(text)
            list_paragraph = template.add_paragraph()
            apply_paragraph_format(list_paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY, spacing)
            apply_numbering(list_paragraph, current_num_pr)
            add_run(list_paragraph, item_text)
            continue

        in_list = False

        if first_paragraph:
            add_paragraph_with_runs(
                template,
                [(text, True)],
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                spacing=STANDARD_SPACE,
            )
            first_paragraph = False
            continue

        if is_slide_heading(text):
            label = normalize_slide_label(text)
            add_paragraph_with_runs(
                template,
                [(label, True)],
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                spacing=STANDARD_SPACE,
            )
            continue

        if text.startswith("Доклад окончен"):
            add_paragraph_with_runs(
                template,
                [(text, True)],
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                spacing=STANDARD_SPACE,
            )
            continue

        if text == "Рекомендации для ответов на вопросы":
            template.add_page_break()
            add_paragraph_with_runs(
                template,
                [(text, True)],
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                spacing=STANDARD_SPACE,
            )
            continue

        match = QUESTION_RE.match(text)
        if match:
            label = match.group(1)
            rest = match.group(2)
            paragraph_out = template.add_paragraph()
            apply_paragraph_format(paragraph_out, WD_ALIGN_PARAGRAPH.JUSTIFY, STANDARD_SPACE)
            label_text = f"{label}:"
            if rest:
                label_text += " "
            add_run(paragraph_out, label_text, bold=True)
            if rest:
                add_run(paragraph_out, rest)
            continue

        split_result = split_schema_or_poster(text)
        if split_result is not None:
            lead, tail = split_result
            runs = [(lead, True)]
            if tail:
                runs.append((f" {tail}", False))
            add_paragraph_with_runs(
                template,
                runs,
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                spacing=STANDARD_SPACE,
            )
            continue

        next_paragraph = next_non_empty(source_doc.paragraphs, idx)
        spacing = STANDARD_SPACE
        if text.endswith(":") and next_paragraph is not None and is_list_item(next_paragraph):
            spacing = None

        if is_all_bold(paragraph):
            add_paragraph_with_runs(
                template,
                [(text, True)],
                alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
                spacing=spacing,
            )
            continue

        add_runs_from_source(template, paragraph, WD_ALIGN_PARAGRAPH.JUSTIFY, spacing)

    template.save(str(output_path))
    return output_path


def main() -> None:
    base_dir = Path(__file__).resolve().parent
    template_path = base_dir / TEMPLATE_FILENAME
    source_path = Path(sys.argv[1]) if len(sys.argv) > 1 else base_dir / DEFAULT_INPUT_FILENAME
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else base_dir / DEFAULT_OUTPUT_FILENAME

    format_document(source_path, output_path, template_path)


if __name__ == "__main__":
    main()
